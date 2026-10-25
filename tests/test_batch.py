"""Fase 14: planillas, proyección de redes a buques y cálculo por lotes."""

from __future__ import annotations

import csv
import io
import math
from functools import cache

import pytest

from trawlkit import DesignCase, run
from trawlkit.batch import (
    estimate_cost_s,
    fleet_vessel,
    project_net,
    read_book,
    read_input,
    read_nets,
    result_xlsx,
    run_input,
    template,
    to_csv,
)
from trawlkit.batch.columns import FLEET
from trawlkit.batch.nets import PANEL_HEADERS, net_rows
from trawlkit.batch.project import ProjectOptions
from trawlkit.batch.sheets import parse_sheet, read_csv
from trawlkit.io import load_doors, load_net, load_vessels

FAST = ProjectOptions(assemble=False)


def _a3_csv() -> str:
    spec = load_net("fao_a3_standard")
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([h for _, h in PANEL_HEADERS])
    for r in net_rows(spec):
        w.writerow([str(x).replace(".", ",") if isinstance(x, float) else x for x in r])
    return buf.getvalue()


def test_fleet_csv_with_aliases_units_and_comma_decimals() -> None:
    text = (
        "# flota del puerto\n"
        "Nombre;Potencia (kW);Eslora [ft];Hélice;Profundidad;Región;Fondo;Solo genéricos\n"
        "Don Pedro;44,1;36;PF;30;Latinoamérica;arena;sí\n"
        ";;;;;;;\n"
        "Sin potencia;;;;;;;\n"
        "Rara;abc;;PX;;;;\n"
    )
    p = parse_sheet(read_csv(text), FLEET)
    assert p.columns["Potencia (kW)"] == "bhp"
    rec = p.records[0]
    assert rec["bhp"] == pytest.approx(44.1 / 0.7355)
    assert rec["length_m"] == pytest.approx(36 * 0.3048)
    assert rec["region"] == "latam" and rec["seabed"] == "sand" and rec["generic_only"] == "true"
    codes = {(i.row, i.code) for i in p.issues}
    assert (2, "missing_value") in codes and (3, "bad_value") in codes
    # la fila vacía no cuenta; los errores no abortan las demás filas
    assert len(p.records) == 3
    fv = fleet_vessel(rec)
    assert fv.vessel.bhp == pytest.approx(60, rel=1e-3) and fv.generic_only
    assert fv.vessel.beam_m and any("beam_m" in n for n in fv.notes)
    unit = parse_sheet(read_csv("nombre,potencia\nX,110 kW\n"), FLEET).records[0]
    assert unit["bhp"] == pytest.approx(110 / 0.7355)


def test_a3_net_round_trip_through_a_sheet_gives_the_same_result() -> None:
    spec = load_net("fao_a3_standard")
    nets, issues = read_nets([read_csv(_a3_csv(), "paneles")])
    assert issues == [] and len(nets) == 1
    loaded = nets[0]
    assert loaded.ok and loaded.spec.status == "published"
    assert [p.id for p in loaded.spec.panels] == [p.id for p in spec.panels]
    v = next(x for x in load_vessels().vessels if x.id == "fao_a3_150hp")
    d = load_doors().doors[0]
    full = loaded.spec.model_copy(
        update={k: getattr(spec, k) for k in ("species", "headrope", "footrope", "floats")}
    )
    a = run(DesignCase(spec, v, d)).summary()
    b = run(DesignCase(full, v, d)).summary()
    for k, x in a.items():
        if isinstance(x, float):
            assert b[k] == pytest.approx(x, rel=1e-12), k


def test_broken_seam_in_a_sheet_is_flagged_as_draft() -> None:
    text = _a3_csv().splitlines()
    cells = text[3].split(";")  # panel C: borde de bolsa roto
    cells[9] = "300"
    text[3] = ";".join(cells)
    nets, _ = read_nets([read_csv("\n".join(text), "paneles")])
    assert not nets[0].ok and nets[0].spec.status == "draft"
    assert any(i.level == "error" for i in nets[0].check)


@cache
def _projection(net: str, hp: float):  # type: ignore[no-untyped-def]
    return project_net(net, fleet_vessel({"name": f"B{hp:g}", "bhp": hp}), FAST)


def test_projection_follows_the_power_similarity_law() -> None:
    same = _projection("fao_a3_standard", 150.0)
    assert same.s0 == pytest.approx(1.0)
    assert 0.8 <= same.s <= 1.05 and same.feasible
    # La escala final sigue a la de semejanza: s/s₀ casi constante entre 40 y 300 hp.
    ratios = [
        _projection("fao_wb_bt2p_60hp", hp).s / math.sqrt(hp / 60) for hp in (40.0, 150.0, 300.0)
    ]
    assert max(ratios) - min(ratios) < 0.1
    big = _projection("fao_wb_bt2p_60hp", 150.0)
    assert big.s0 == pytest.approx(2.5**0.5)
    assert big.feasible
    assert 0.7 * big.s0 <= big.s <= 1.3 * big.s0
    small = _projection("fao_wb_bt2p_60hp", 30.0)
    assert small.s < big.s
    # la red escalada conserva mallero e hilo y escala las mallas
    base = load_net("fao_wb_bt2p_60hp")
    p0, p1 = base.panels[0], big.net.panels[0]
    assert p1.mesh_mm == p0.mesh_mm and p1.meshes_depth > p0.meshes_depth
    row = big.summary()
    assert row["s0"] < 1.6 and row["feasible"] and row["warp"]


def test_projection_with_assembly_and_fixed_speed() -> None:
    fv = fleet_vessel({"name": "Rápido", "bhp": 150, "speed_kn": 3.0, "region": "latam"})
    p = project_net("fao_wb_bt2p_60hp", fv)
    assert p.assembly is not None and p.summary()["speed_kn"] == pytest.approx(3.0)
    assert p.case is not None and p.case.net.id == p.net.id
    assert p.case.warp_id == p.assembly.warp_id


def test_templates_round_trip_and_batch_export() -> None:
    for kind in ("fleet", "cases", "nets"):
        for fmt in ("xlsx", "csv"):
            inp = read_input(template(kind, fmt), f"t.{fmt}", kind)  # type: ignore[arg-type]
            assert inp.size > 0
            assert [i for i in inp.issues if i.level == "error"] == [], (kind, fmt)
    inp = read_input(template("fleet", "xlsx"), "flota.xlsx", "fleet")
    res = run_input(inp, net="fao_wb_bt2p_60hp", options=ProjectOptions(max_doors=2))
    assert len(res.rows) == 3 and res.counts["error"] == 0
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(result_xlsx(res)))
    assert wb.sheetnames[:2] == ["resumen", "materiales"]
    assert any(n.startswith("red ") for n in wb.sheetnames)
    assert wb["materiales"].max_row > 3
    csv_text = to_csv(res)
    assert csv_text.startswith("Fila;Nombre;Estado;Buque") and "cumple" in csv_text
    cases = run_input(read_input(template("cases", "csv"), "c.csv", "cases"))
    assert cases.counts == {"ok": 3, "infeasible": 0, "error": 0}
    assert cases.rows[1].summary["vessel_id"] == "fao_a3_180hp"


def test_cases_errors_do_not_stop_the_batch() -> None:
    text = "caso;red;buque;sweeps_m;nada\nbien;fao_a3_standard;;30;\nmal;no_existe;;;\n"
    inp = read_input(text, "c.csv", "cases")
    assert any(i.code == "unknown_column" for i in inp.issues)
    res = run_input(inp)
    assert [r.status for r in res.rows] == ["ok", "error"]
    assert res.rows[0].summary["net_id"] == "fao_a3_standard"


def test_cost_estimate_and_sheet_kinds() -> None:
    assert estimate_cost_s("fleet", 10) > estimate_cost_s("cases", 10)
    sheets = read_book(template("nets", "xlsx"), "redes.xlsx")
    assert [s.name for s in sheets] == ["paneles", "red", "ayuda"]
    empty = read_input("nombre;potencia\n", "f.csv", "fleet")
    assert any(i.code == "empty" for i in empty.issues)


def test_cli_batch_template_run_and_pdf(tmp_path) -> None:  # type: ignore[no-untyped-def]
    from typer.testing import CliRunner

    from trawlkit.cli import app

    runner = CliRunner()
    tpl = tmp_path / "flota.xlsx"
    r = runner.invoke(app, ["batch", "--template", str(tpl)])
    assert r.exit_code == 0 and tpl.exists()
    out = tmp_path / "lote.xlsx"
    r = runner.invoke(
        app, ["batch", str(tpl), "--net", "fao_wb_bt2p_60hp", "--no-assemble", "-o", str(out)]
    )
    assert r.exit_code == 0, r.output
    assert "3 filas" in r.output and out.stat().st_size > 1000
    pytest.importorskip("weasyprint")
    pdf = tmp_path / "lote.pdf"
    r = runner.invoke(
        app, ["batch", str(tpl), "--net", "fao_wb_bt2p_60hp", "--no-assemble", "-o", str(pdf)]
    )
    assert r.exit_code == 0, r.output
    assert pdf.read_bytes()[:4] == b"%PDF"
