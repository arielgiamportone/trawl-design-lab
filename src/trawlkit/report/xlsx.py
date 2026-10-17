"""Libro XLSX del informe con openpyxl (extra ``report`` o ``excel``).

Una hoja por :class:`~trawlkit.report.tables.Table` (nombre = ``Table.key`` recortado a 31
caracteres), encabezados en negrita con unidad, formatos numéricos, anchos de columna y paneles
inmovilizados; más una hoja "Caso" con ``DesignCase.to_dict()`` aplanado y una hoja "Notas".
"""

from __future__ import annotations

from collections.abc import Iterator, Sequence
from pathlib import Path
from typing import TYPE_CHECKING, Any

from trawlkit.report.tables import Table, all_tables

if TYPE_CHECKING:
    from trawlkit.model import DesignResult

#: Largo máximo del nombre de hoja en Excel.
MAX_SHEET_NAME = 31


def _openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise RuntimeError(
            "La salida XLSX necesita openpyxl: pip install trawlkit[report]"
        ) from exc
    return openpyxl


def sheet_name(key: str, used: set[str]) -> str:
    """Nombre de hoja válido para Excel (≤ 31 caracteres, sin caracteres prohibidos), único."""
    clean = "".join("_" if ch in "[]:*?/\\" else ch for ch in key)[:MAX_SHEET_NAME] or "hoja"
    name, n = clean, 2
    while name.lower() in {u.lower() for u in used}:
        suffix = f"_{n}"
        name = clean[: MAX_SHEET_NAME - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def number_format(value: float) -> str:
    """``0.0000`` para valores pequeños (|x| < 1), ``0`` para enteros y ``0.00`` en general."""
    if float(value).is_integer() and abs(value) >= 1:
        return "0"
    return "0.0000" if abs(value) < 1 else "0.00"


def flatten(data: Any, prefix: str = "") -> Iterator[tuple[str, Any]]:
    """Aplana diccionarios y listas en pares ``(clave.punteada, valor escalar)``."""
    if isinstance(data, dict):
        if not data and prefix:
            yield prefix, ""
        for k, v in data.items():
            yield from flatten(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(data, list | tuple):
        if all(not isinstance(x, dict | list | tuple) for x in data):
            yield prefix, ", ".join(str(x) for x in data)
        else:
            for i, v in enumerate(data):
                yield from flatten(v, f"{prefix}[{i}]")
    else:
        yield prefix, data


def _write_table(ws: Any, table: Table, font: Any) -> None:
    from openpyxl.utils import get_column_letter

    ws.append([table.title])
    ws.cell(row=1, column=1).font = font(bold=True, size=12)
    ws.append([c.heading for c in table.columns])
    for cell in ws[2]:
        cell.font = font(bold=True)
    widths = [len(c.heading) for c in table.columns]
    for row in table.rows:
        ws.append(list(row))
        r = ws.max_row
        for j, value in enumerate(row):
            if isinstance(value, float):
                ws.cell(row=r, column=j + 1).number_format = number_format(value)
            shown = f"{value:.4g}" if isinstance(value, float) else str(value)
            widths[j] = max(widths[j], len(shown))
    extra = ws.max_row + 2
    if table.note:
        ws.cell(row=extra, column=1, value=table.note)
        extra += 1
    if table.source:
        ws.cell(row=extra, column=1, value=f"Celdas del libro: {table.source}")
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 8), 60)
    ws.freeze_panes = "B3"


def write_xlsx(result: DesignResult, path: str | Path, extra_tables: Sequence[Table] = ()) -> Path:
    """Escribe el libro XLSX del informe en ``path`` y devuelve la ruta.

    ``extra_tables`` (p. ej. las de optimización y simulación de :mod:`trawlkit.report.extras`)
    van como hojas adicionales antes de "Caso" y "Notas".
    """
    openpyxl = _openpyxl()
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    used: set[str] = set()
    first = True
    # La tabla de notas va en la hoja "Notas" (al final, junto a "Caso").
    tables = [t for t in all_tables(result) if t.key != "notas"] + list(extra_tables)
    for table in tables:
        if first:
            ws = wb.active
            ws.title = sheet_name(table.key, used)
            first = False
        else:
            ws = wb.create_sheet(sheet_name(table.key, used))
        _write_table(ws, table, Font)

    ws = wb.create_sheet(sheet_name("Caso", used))
    ws.append(["Parámetro", "Valor"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for key, value in flatten(result.case.to_dict()):
        ws.append(
            [key, value if isinstance(value, int | float | str) or value is None else str(value)]
        )
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 40
    ws.freeze_panes = "A2"

    ws = wb.create_sheet(sheet_name("Notas", used))
    ws.append(["N.º", "Nota"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, note in enumerate(result.notes, 1):
        ws.append([i, note])
    ws.column_dimensions["B"].width = 120
    ws.freeze_panes = "A2"

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


__all__ = ["flatten", "number_format", "sheet_name", "write_xlsx"]
