"""Exportación del lote (CSV y XLSX) y plantillas para cargar planillas (Fase 14).

El XLSX de un lote de flota tiene las hojas «resumen» (una fila por buque), «materiales» (la
lista de materiales de cada armado), «observaciones» y una hoja de paneles por red escalada,
lista para volver a cargarla. El CSV es la hoja de resumen, con ``;`` y coma decimal (abre bien
en planillas en castellano).
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from typing import Any

from trawlkit.batch.columns import CASES, FLEET, NETS, PANELS
from trawlkit.batch.nets import PANEL_HEADERS, net_rows
from trawlkit.batch.run import BatchResult
from trawlkit.batch.sheets import Col
from trawlkit.io import load_net

Table = tuple[str, list[str], list[list[Any]]]

#: Columnas del resumen por modo: (clave de ``BatchRow.summary``, encabezado).
SUMMARY_COLUMNS: dict[str, tuple[tuple[str, str], ...]] = {
    "fleet": (
        ("vessel", "Buque"),
        ("bhp", "Potencia [hp]"),
        ("base_net", "Red base"),
        ("s0", "Escala por semejanza"),
        ("s", "Escala"),
        ("speed_kn", "Velocidad [kn]"),
        ("feasible", "Cumple"),
        ("headrope_m", "Lima [m]"),
        ("footrope_m", "Burlón [m]"),
        ("ah_m", "Abertura horizontal [m]"),
        ("av_m", "Abertura vertical [m]"),
        ("demand_kgf", "Resistencia total [kgf]"),
        ("thrust_margin", "Reserva de tiro"),
        ("area_rate_m2_h", "Área barrida [m²/h]"),
        ("door", "Portón"),
        ("door_area_m2", "Superficie del portón [m²]"),
        ("warp", "Cable"),
        ("float", "Boya"),
        ("float_count", "Boyas"),
        ("twine", "Hilo"),
        ("cost", "Costo orientativo"),
        ("currency", "Moneda"),
        ("missing_prices", "Líneas sin precio"),
        ("failed", "No cumple"),
    ),
    "cases": (
        ("case", "Caso"),
        ("net_id", "Red"),
        ("vessel_id", "Buque"),
        ("door_id", "Portón"),
        ("headrope_m", "Lima [m]"),
        ("footrope_m", "Burlón [m]"),
        ("vertical_opening_m", "Abertura vertical [m]"),
        ("horizontal_opening_m", "Abertura horizontal [m]"),
        ("rn_std_3kn_kgf", "Resistencia de la red a 3 kn [kgf]"),
        ("max_speed_kn", "Velocidad máxima [kn]"),
        ("warp_selected", "Cable"),
        ("float_count", "Boyas"),
        ("net_weight_kg", "Peso de la red [kg]"),
    ),
    "nets": (
        ("net_id", "Red"),
        ("panels", "Paneles"),
        ("errors", "Errores del validador"),
        ("warnings", "Avisos"),
        ("headrope_m", "Lima [m]"),
        ("footrope_m", "Burlón [m]"),
        ("vertical_opening_m", "Abertura vertical [m]"),
        ("horizontal_opening_m", "Abertura horizontal [m]"),
        ("rn_std_3kn_kgf", "Resistencia de la red a 3 kn [kgf]"),
        ("max_speed_kn", "Velocidad máxima [kn]"),
    ),
}

#: Columnas económicas del lote de flota (Fase 16), si se pidió la economía.
ECONOMIC_COLUMNS: tuple[tuple[str, str], ...] = (
    ("fuel_l_trip", "Gasoil por marea [L]"),
    ("catch_kg_trip", "Captura por marea [kg]"),
    ("revenue_trip", "Ingreso por marea"),
    ("cost_trip", "Costo por marea"),
    ("cost_per_kg", "Costo por kg"),
    ("margin_trip", "Margen por marea"),
    ("margin_year", "Margen por año"),
)

STATUS_LABEL = {"ok": "cumple", "infeasible": "no cumple", "error": "error"}


def _cell(value: Any) -> Any:
    if isinstance(value, bool):
        return "sí" if value else "no"
    if isinstance(value, list | tuple):
        return "; ".join(str(x) for x in value)
    if isinstance(value, float):
        return round(value, 4)
    return "" if value is None else value


def tables(result: BatchResult) -> list[Table]:
    """Hojas del lote: resumen, materiales (flota), redes escaladas y observaciones."""
    cols = SUMMARY_COLUMNS[result.mode]
    if result.mode == "fleet" and any("fuel_l_trip" in r.summary for r in result.rows):
        cols = cols[:-1] + ECONOMIC_COLUMNS + cols[-1:]
    head = ["Fila", "Nombre", "Estado"] + [h for _, h in cols]
    rows = [
        [r.index + 1, r.name, STATUS_LABEL[r.status], *(_cell(r.summary.get(k)) for k, _ in cols)]
        if r.status != "error"
        else [r.index + 1, r.name, "error", r.error or ""]
        for r in result.rows
    ]
    out: list[Table] = [("resumen", head, rows)]
    if result.mode == "fleet":
        mat: list[list[Any]] = []
        for r in result.rows:
            p = r.projection
            if p is None or p.assembly is None:
                continue
            for ln in p.assembly.bom.lines:
                mat.append(
                    [
                        r.name,
                        ln.label,
                        ln.component_name or ln.component_id or "",
                        round(ln.quantity, 3),
                        ln.unit,
                        _cell(ln.unit_price),
                        ln.currency or "",
                        "sin precio" if ln.cost is None else ln.cost,
                    ]
                )
        out.append(
            (
                "materiales",
                [
                    "Buque",
                    "Ítem",
                    "Componente",
                    "Cantidad",
                    "Unidad",
                    "Precio unitario",
                    "Moneda",
                    "Costo",
                ],
                mat,
            )
        )
        for r in result.rows:
            if r.projection is not None:
                out.append(
                    (f"red {r.name}", [h for _, h in PANEL_HEADERS], net_rows(r.projection.net))
                )
    notes = [[r.index + 1, r.name, n] for r in result.rows for n in r.notes]
    notes += [
        [
            i.row or "",
            i.sheet or "",
            f"{i.level}: {i.message}" + (f" ({i.column})" if i.column else ""),
        ]
        for i in result.issues
    ]
    out.append(("observaciones", ["Fila", "Buque, caso o red", "Observación"], notes))
    return out


def to_csv(result: BatchResult) -> str:
    """Hoja de resumen en CSV (``;``, coma decimal)."""
    name, head, rows = tables(result)[0]
    return _csv(head, rows)


def _csv(head: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    w.writerow(head)
    for row in rows:
        w.writerow([str(v).replace(".", ",") if isinstance(v, float) else v for v in row])
    return buf.getvalue()


def to_xlsx(sheets: Sequence[Table]) -> bytes:
    """Libro XLSX con las hojas dadas (encabezados en negrita, anchos y paneles fijos)."""
    try:
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise RuntimeError('el XLSX necesita openpyxl: pip install "trawlkit[excel]"') from exc
    from trawlkit.report.xlsx import sheet_name

    wb = openpyxl.Workbook()
    used: set[str] = set()
    for i, (name, head, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = sheet_name(name, used)
        ws.append(list(head))
        for c in ws[1]:
            c.font = Font(bold=True)
        widths = [len(str(h)) for h in head]
        for row in rows:
            ws.append(list(row))
            for j, v in enumerate(row[: len(widths)]):
                widths[j] = max(widths[j], len(f"{v:.4g}" if isinstance(v, float) else str(v)))
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 8), 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def result_xlsx(result: BatchResult) -> bytes:
    """XLSX del lote (ver :func:`tables`)."""
    return to_xlsx(tables(result))


# -- plantillas --------------------------------------------------------------------------------

TemplateKind = str  # "fleet" | "nets" | "cases"

_EXAMPLE_FLEET: list[list[Any]] = [
    ["Don Pedro", 60, 11, 3.6, 1.6, "PF", 30, "", "latam", "arena", "sí"],
    ["La Gaviota", "110 kW", 14, 4.5, 2.0, "PF", 50, 2.5, "latam", "fango", "no"],
    ["Estrella", 250, "", "", "", "PV+T", 80, "", "", "mixto", ""],
]

_EXAMPLE_CASES: list[list[Any]] = [
    ["A-3 libro", "fao_a3_standard", "fao_a3_150hp", "fao_flat_wood_a3", "corregido", 25, ""],
    ["A-3 180 hp", "fao_a3_standard", "fao_a3_180hp", "fao_flat_wood_a3", "corregido", 25, 60],
    ["Workbook 60 hp", "fao_wb_bt2p_60hp", "", "", "corregido", "", ""],
]


def _help_rows(cols: Sequence[Col]) -> list[list[Any]]:
    out = []
    for c in cols:
        rng = "" if c.low is None else f"{c.low:g}–{c.high:g}"
        out.append(
            [
                c.key,
                ", ".join(c.aliases[:4]),
                c.unit,
                "sí" if c.required else "",
                c.help,
                c.example,
                rng,
            ]
        )
    return out


HELP_HEAD = ["Columna", "También se acepta", "Unidad", "Obligatoria", "Qué es", "Ejemplo", "Rango"]


def template_tables(kind: TemplateKind) -> list[Table]:
    """Hojas de la plantilla de ``kind`` (``fleet``, ``nets`` o ``cases``) con ejemplos y ayuda."""
    if kind == "fleet":
        head = [
            "nombre",
            "potencia [hp]",
            "eslora [m]",
            "manga [m]",
            "puntal [m]",
            "hélice",
            "profundidad [m]",
            "velocidad [kn]",
            "región",
            "fondo",
            "sólo genéricos",
        ]
        return [("flota", head, _EXAMPLE_FLEET), ("ayuda", HELP_HEAD, _help_rows(FLEET))]
    if kind == "nets":
        spec = load_net("fao_wb_bt2p_60hp")
        rows = [[*r[:1], *r[1:]] for r in net_rows(spec)]
        for r in rows:
            r[0] = "Red de ejemplo (Workbook 60 hp)"
        net_head = [
            "red",
            "descripción",
            "familia",
            "aparejo",
            "hp_min",
            "hp_max",
            "profundidad [m]",
        ]
        net_row = [
            [
                "Red de ejemplo (Workbook 60 hp)",
                "Copia de la red del Workbook",
                "fondo",
                "portones",
                50,
                70,
                30,
            ]
        ]
        return [
            ("paneles", [h for _, h in PANEL_HEADERS], rows),
            ("red", net_head, net_row),
            ("ayuda", HELP_HEAD, _help_rows(PANELS) + _help_rows(NETS)),
        ]
    if kind == "cases":
        head = ["caso", "red", "buque", "portón", "modo", "sweeps_m", "depth_m"]
        help_rows: list[list[Any]] = _help_rows(CASES) + [
            [
                "(cualquier parámetro)",
                "",
                "",
                "",
                "columnas con el nombre de GET /params (sweeps_m, depth_m, door_attack_angle_deg…)",
                "25",
                "",
            ]
        ]
        return [("casos", head, _EXAMPLE_CASES), ("ayuda", HELP_HEAD, help_rows)]
    raise ValueError(f"plantilla desconocida: {kind!r} (fleet, nets o cases)")


def template(kind: TemplateKind, fmt: str = "xlsx") -> bytes | str:
    """Plantilla en XLSX (con hoja de ayuda) o CSV (la primera hoja)."""
    sheets = template_tables(kind)
    if fmt == "csv":
        name, head, rows = sheets[0]
        return _csv(head, rows)
    if fmt != "xlsx":
        raise ValueError("formato: xlsx o csv")
    return to_xlsx(sheets)


__all__ = [
    "SUMMARY_COLUMNS",
    "result_xlsx",
    "tables",
    "template",
    "template_tables",
    "to_csv",
    "to_xlsx",
]
